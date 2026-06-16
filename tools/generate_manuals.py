from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "docs" / "manuales"
DATE_TEXT = datetime.now().strftime("%d/%m/%Y")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 99, 110)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


STATES = [
    ("Arribo al pais", "Estado inicial luego de crear la guia madre."),
    ("Arribo a aduana", "Estado posterior a entregar la guia madre a aduana."),
    ("Carga de selectivos", "Estado posterior a cargar consolidados/selectivos."),
    ("Completado", "Estado final luego de despachar todas las cajas habilitadas."),
]

FLOW_STEPS = [
    ("1", "Login", "Seleccionar pais e iniciar sesion con credenciales autorizadas.", "Pantalla Guias madre visible."),
    ("2", "Agregar guia madre", "Completar datos obligatorios y crear la guia.", "Guia en Arribo al pais."),
    ("3", "Entrega a aduana", "Usar Entregar a aduana sobre la guia creada.", "Guia en Arribo a aduana."),
    ("4", "Agregar consolidado", "Entrar al detalle y cargar Verde, Rojo y Amarillo.", "Guia en Carga de selectivos."),
    ("5", "Despacho de cajas", "Despachar todos los consolidados habilitados.", "Guia en Completado."),
]

COUNTRIES = ["Guatemala", "Honduras", "El Salvador"]
SELECTIVES = [
    ("Verde", "Archivo_de_prueba_con_una_hoja_verde.xlsx", "Prueba de automatizacion_1"),
    ("Rojo", "Archivo_de_prueba_con_una_hoja_rojo.xlsx", "Prueba de automatizacion_2"),
    ("Amarillo", "Archivo_de_prueba_con_una_hoja_amarillo.xlsx", "Prueba de automatizacion_3"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def apply_styles(doc: Document, title: str, subtitle: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = title
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.text = "Control Aduanero FD | Documento generado para QA"
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph()
    run = p.add_run(subtitle)
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(12)

    add_info_table(doc, [
        ("Proyecto", "Control Aduanero FD"),
        ("Fecha", DATE_TEXT),
        ("Version", "1.0"),
        ("Base", "Flujos automatizados y validados en QA"),
    ])


def add_info_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.65, 4.85])
    for idx, (label, value) in enumerate(rows):
        table.cell(idx, 0).text = label
        table.cell(idx, 1).text = value
        set_cell_shading(table.cell(idx, 0), HEADER_FILL)
        for cell in table.row_cells(idx):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
            if idx == 0:
                pass
    doc.add_paragraph()


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, HEADER_FILL)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = INK
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = value
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
    set_table_width(table, widths)
    doc.add_paragraph()


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.size = Pt(10)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_common_flow_sections(doc: Document, technical: bool = False) -> None:
    doc.add_heading("Flujo operativo validado", level=1)
    add_matrix(
        doc,
        ["Paso", "Modulo / accion", "Actividad", "Estado esperado"],
        FLOW_STEPS,
        [0.55, 1.65, 2.85, 1.45],
    )

    doc.add_heading("Estados principales", level=1)
    add_matrix(
        doc,
        ["Estado", "Descripcion"],
        [[state, desc] for state, desc in STATES],
        [1.7, 4.8],
    )

    doc.add_heading("Paises y alcance", level=1)
    add_bullets(doc, [f"El flujo aplica para {country}." for country in COUNTRIES])

    doc.add_heading("Archivos de consolidado", level=1)
    add_matrix(
        doc,
        ["Selectivo", "Archivo", "Nombre sugerido"],
        [[selective, file_name, name] for selective, file_name, name in SELECTIVES],
        [1.2, 3.35, 1.95],
    )

    if technical:
        doc.add_heading("Marcadores Pytest", level=1)
        add_matrix(
            doc,
            ["Marcador", "Proposito"],
            [
                ["login", "Login exitoso por pais."],
                ["login_negativo", "Credenciales invalidas."],
                ["guia_madre", "Creacion de guia madre por pais y moneda."],
                ["guia_madre_negativo", "Campos obligatorios de guia madre."],
                ["entrega_aduana", "Cambio de Arribo al pais a Arribo a aduana."],
                ["agregar_consolidado", "Carga de selectivos Verde, Rojo y Amarillo."],
                ["despacho_cajas", "Despacho de consolidados habilitados."],
                ["flujo_completo", "End to end: login hasta despacho."],
            ],
            [1.85, 4.65],
        )


def build_word_general() -> None:
    doc = Document()
    apply_styles(
        doc,
        "Documento Word - Flujo Control Aduanero FD",
        "Resumen operativo del proceso automatizado de punta a punta.",
    )
    add_callout(
        doc,
        "Objetivo",
        "Documentar en un solo documento el flujo completo validado: login, creacion de guia madre, entrega a aduana, carga de consolidados y despacho de cajas.",
    )
    add_common_flow_sections(doc)
    doc.add_heading("Criterios de aceptacion", level=1)
    add_bullets(
        doc,
        [
            "El usuario puede iniciar sesion en el pais seleccionado.",
            "La guia madre creada aparece en el listado y avanza por los estados definidos.",
            "Los consolidados se cargan con el selectivo y archivo correspondiente.",
            "El despacho procesa todos los consolidados habilitados.",
            "Al finalizar, la guia madre queda en estado Completado y no permite agregar nuevos consolidados.",
        ],
    )
    save_doc(doc, "Documento_Word_Flujo_Control_Aduanero_FD.docx")


def build_qa_manual() -> None:
    doc = Document()
    apply_styles(
        doc,
        "Manual tecnico QA - Control Aduanero FD",
        "Guia tecnica para ejecutar, validar y mantener la automatizacion QA.",
    )
    add_callout(
        doc,
        "Uso previsto",
        "Este manual esta orientado a QA Automation y cubre estructura, comandos, evidencias, marcadores, datos y criterios de validacion.",
    )
    doc.add_heading("Estructura del repositorio", level=1)
    add_matrix(
        doc,
        ["Ruta", "Uso"],
        [
            ["features/", "Escenarios BDD en Gherkin."],
            ["tests/", "Steps de Pytest-BDD y generacion de datos."],
            ["pages/", "Page Object Model y acciones de UI."],
            ["docs/manuales/", "Manuales generados para QA y usuario final."],
            ["allure-results/", "Resultados crudos de Allure."],
            ["allure-reports/<marker>/", "Reporte HTML generado por marcador."],
            ["videos/", "Videos WebM por escenario ejecutado."],
        ],
        [2.0, 4.5],
    )
    add_common_flow_sections(doc, technical=True)
    doc.add_heading("Comandos de ejecucion", level=1)
    add_matrix(
        doc,
        ["Comando", "Descripcion"],
        [
            [r".\.venv\Scripts\python.exe -m pytest --collect-only", "Valida que Pytest recolecte los escenarios."],
            [r".\run_login_report.ps1 -Marker login", "Ejecuta login y genera reporte."],
            [r".\run_login_report.ps1 -Marker agregar_consolidado", "Ejecuta carga de consolidados."],
            [r".\run_login_report.ps1 -Marker despacho_cajas", "Ejecuta despacho de cajas."],
            [r".\run_login_report.ps1 -Marker flujo_completo", "Ejecuta el flujo end to end."],
        ],
        [3.4, 3.1],
    )
    doc.add_heading("Variables de entorno", level=1)
    add_matrix(
        doc,
        ["Variable", "Proposito"],
        [
            ["BASE_URL", "URL del ambiente QA."],
            ["CONTROL_ADUANERO_USER", "Usuario autorizado."],
            ["CONTROL_ADUANERO_PASSWORD", "Contrasena del usuario autorizado."],
            ["BROWSER_CHANNEL", "Canal de navegador, por ejemplo msedge."],
            ["HEADLESS", "true/false segun se requiera ver la ejecucion."],
            ["SLOW_MO", "Tiempo en milisegundos para demo visual."],
            ["RECORD_VIDEO", "Activa o desactiva video por escenario."],
            ["CONSOLIDADO_FILES_DIR", "Ruta donde viven los archivos Excel de prueba."],
        ],
        [2.2, 4.3],
    )
    doc.add_heading("Checklist QA antes de demo", level=1)
    add_bullets(
        doc,
        [
            "Confirmar que los archivos Excel de selectivos existan en Descargas o en CONSOLIDADO_FILES_DIR.",
            "Ejecutar el marcador requerido con RECORD_VIDEO=true.",
            "Validar el resumen de Pytest: passed, failed, deselected y warnings.",
            "Abrir el reporte Allure generado en allure-reports/<marker>/index.html.",
            "Confirmar que los videos WebM se generaron en videos/.",
            "Si falla un escenario, revisar la captura y el video adjunto antes de reejecutar.",
        ],
    )
    save_doc(doc, "Manual_Tecnico_QA_Control_Aduanero_FD.docx")


def build_user_manual() -> None:
    doc = Document()
    apply_styles(
        doc,
        "Manual funcional para usuarios finales - Control Aduanero FD",
        "Guia paso a paso para operar el flujo de manifiestos y consolidados.",
    )
    add_callout(
        doc,
        "Audiencia",
        "Este documento esta orientado a usuarios operativos que necesitan ejecutar el flujo dentro del sistema, sin detalles de automatizacion.",
    )
    doc.add_heading("Requisitos previos", level=1)
    add_bullets(
        doc,
        [
            "Contar con usuario y contrasena activos.",
            "Conocer el pais donde se trabajara el manifiesto.",
            "Tener los archivos de consolidado listos para cargar.",
            "Validar que la guia madre tenga la informacion requerida antes de crearla.",
        ],
    )
    doc.add_heading("Procedimiento funcional", level=1)
    add_numbered(
        doc,
        [
            "Ingresar al sistema Control Aduanero FD.",
            "Seleccionar el pais correspondiente e iniciar sesion.",
            "Crear una guia madre completando los campos obligatorios.",
            "Confirmar que la guia quede en Arribo al pais.",
            "Entregar la guia madre a aduana y validar Arribo a aduana.",
            "Entrar al detalle de la guia y cargar los consolidados Verde, Rojo y Amarillo.",
            "Validar que la guia pase a Carga de selectivos.",
            "Entrar al detalle de cada consolidado y despachar las cajas habilitadas.",
            "Confirmar que la guia quede en Completado.",
            "Verificar que el boton Agregar consolidado quede bloqueado.",
        ],
    )
    add_common_flow_sections(doc)
    doc.add_heading("Campos obligatorios de guia madre", level=1)
    add_matrix(
        doc,
        ["Campo", "Descripcion"],
        [
            ["No. Guia", "Identificador de la guia madre."],
            ["Numero de vuelo", "Vuelo asociado al manifiesto."],
            ["Origen", "Pais o estacion origen."],
            ["Destino", "Pais o estacion destino."],
            ["Total cajas", "Cantidad total de cajas."],
            ["Total de guias individuales", "Cantidad de guias individuales."],
            ["Peso declarado en KG", "Peso total declarado."],
            ["Valor declarado", "Valor monetario declarado."],
        ],
        [2.2, 4.3],
    )
    doc.add_heading("Errores frecuentes", level=1)
    add_matrix(
        doc,
        ["Situacion", "Accion recomendada"],
        [
            ["No se puede crear la guia madre", "Revisar que todos los campos obligatorios esten completos."],
            ["No aparece Entregar a aduana", "Confirmar que la guia este en Arribo al pais."],
            ["No permite agregar consolidado", "Confirmar que la guia este en Arribo a aduana y no Completado."],
            ["No permite despachar", "Validar que el consolidado tenga cajas pendientes de despacho."],
            ["La guia no queda Completado", "Despachar todos los consolidados habilitados de la guia."],
        ],
        [2.35, 4.15],
    )
    save_doc(doc, "Manual_Funcional_Usuario_Final_Control_Aduanero_FD.docx")


def save_doc(doc: Document, filename: str) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    doc.save(path)
    print(path)


def main() -> None:
    build_word_general()
    build_qa_manual()
    build_user_manual()


if __name__ == "__main__":
    main()
